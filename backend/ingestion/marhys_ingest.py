# SPDX-License-Identifier: AGPL-3.0-or-later
# Based on Abyssal Claims — © 2026 Michal Mazurowski — https://something-rare.com
"""Parse the MARHYS 4.0 workbook into rows, passing the source through unchanged.

Source: MARHYS Database 4.0 — Diehl, Alexander; Bach, Wolfgang (2024).
        PANGAEA, https://doi.org/10.1594/PANGAEA.972999   CC-BY-4.0
⭐ The publisher requires the base publication to be cited ALONGSIDE the dataset:
   Diehl & Bach (2020), https://doi.org/10.1029/2020GC009385

⚠️ The DOI's `format=textfile` distribution is NOT the data — it is a 6 KB manifest
listing six binary attachments. The data exists only inside `MARHYS_DB_4_0.xlsx`.

Workbook layout, confirmed by inspection on 2026-09-22 (never from memory):
    sheet   "Tabelle1", the only sheet
    row 12  column headers, 204 columns
    row 13  units
    row 14  first data row
    6788 rows carry a Sample-ID.
"""

from __future__ import annotations

import io
import math
from typing import Any

import httpx
import openpyxl

MARHYS_XLSX_URL = "https://download.pangaea.de/dataset/972999/files/MARHYS_DB_4_0.xlsx"
MARHYS_DOI = "https://doi.org/10.1594/PANGAEA.972999"
MARHYS_BASE_PUBLICATION_DOI = "https://doi.org/10.1029/2020GC009385"
MARHYS_VERSION = "4.0"

SHEET = "Tabelle1"
HEADER_ROW = 12
UNITS_ROW = 13
FIRST_DATA_ROW = 14

#: Column indices are 0-based into the tuple openpyxl yields for a row.
_ID = 1

#: Text/metadata columns → their database field name.
_TEXT_COLUMNS: dict[int, str] = {
    1: "sample_id", 2: "vent_id", 3: "vent_site", 4: "vent_area",
    5: "volcanic_edifice", 6: "rock_type_primary", 7: "rock_type_secondary",
    8: "region_small", 9: "region_large", 10: "geologic_setting",
    11: "expedition", 12: "vessel", 13: "rov_dsv", 14: "dive_no",
    15: "sampler_type", 16: "date_raw", 17: "author_primary",
    19: "author_secondary", 23: "description_a", 24: "description_b",
    25: "sample_type",
}

#: ⛔ The source writes two DIFFERENT placeholders and they do not mean the same
#: thing: "not given" is unknown, "not applicable" is a field that cannot apply
#: to this sample. Both become NULL here — we have nowhere honest to put the
#: distinction — but they are never silently merged into an empty string that
#: would read as a real, blank value. `sampler_type` alone is "not given" on
#: 930 rows (13.7%), so this is not a rare path.
_PLACEHOLDERS = {"not given", "not applicable", "n/a", "na", "-", "?", ""}

#: Numeric columns that get a column of their own, chosen by measured fill rate
#: on 2026-09-22 (every one of these is ≥12% populated across 6788 rows) plus
#: the dissolved gases, which carry the layer's scientific point even where
#: sparser. Everything else lands in `params`.
#:
#: ⭐ THE UNIT IS PART OF THE NAME, deliberately. A reader who sees `fe_umol_kg`
#: cannot mistake it for mg/kg. The source declares these units in row 13 and
#: `parse_marhys` asserts they still match, so a republished file that changed a
#: unit fails loudly instead of silently rescaling every value by 1000.
_NUMERIC_COLUMNS: dict[int, tuple[str, str]] = {
    26: ("lat", "dec°"),
    27: ("lon", "dec°"),
    28: ("depth_mbsl", "mbsl"),
    32: ("temp_c", "°C"),
    36: ("ph", "[]"),
    37: ("alkalinity_mmol_kg", "mmol/kg"),
    38: ("salinity_g_kg", "g/kg"),
    47: ("b_umol_kg", "µmol/kg"),
    48: ("ba_umol_kg", "µmol/kg"),
    51: ("ca_mmol_kg", "mmol/kg"),
    54: ("cu_umol_kg", "µmol/kg"),
    56: ("cs_nmol_kg", "nmol/kg"),
    57: ("fe_umol_kg", "µmol/kg"),
    62: ("k_mmol_kg", "mmol/kg"),
    63: ("li_umol_kg", "µmol/kg"),
    64: ("mg_mmol_kg", "mmol/kg"),
    65: ("mn_umol_kg", "µmol/kg"),
    69: ("na_mmol_kg", "mmol/kg"),
    70: ("nh3_mmol_kg", "mmol/kg"),
    74: ("rb_umol_kg", "µmol/kg"),
    77: ("si_mmol_kg", "mmol/kg"),
    79: ("sr_umol_kg", "µmol/kg"),
    86: ("zn_umol_kg", "µmol/kg"),
    103: ("br_umol_kg", "µmol/kg"),
    104: ("cl_mmol_kg", "mmol/kg"),
    109: ("so4_mmol_kg", "mmol/kg"),
    111: ("ch4_umol_kg", "µmol/kg"),
    113: ("co2_mmol_kg", "mmol/kg"),
    114: ("h2_umol_kg", "µmol/kg"),
    115: ("h2s_mmol_kg", "mmol/kg"),
}

#: Columns that are metadata rather than measurements and must never be swept
#: into `params` as if they were chemistry.
_NON_PARAM_COLUMNS = set(_TEXT_COLUMNS) | set(_NUMERIC_COLUMNS) | {
    18, 20, 21, 22,      # publication years and source ids
}


class MarhysFormatError(RuntimeError):
    """The workbook is not shaped the way this parser was written against."""


def _clean_text(value: Any) -> str | None:
    """Source text, or None for a blank/placeholder. Never an empty string."""
    if value is None:
        return None
    text = str(value).strip()
    return None if text.lower() in _PLACEHOLDERS else text


def _clean_number(value: Any) -> float | None:
    """A real number, or None.

    ⛔ `0.0` IS A REAL VALUE AND IS KEPT. Measured on 2026-09-22: of 1276 zeros in
    the magnesium column, 1252 sit on end-member (`EM`) samples, which have only
    24 non-zero magnesium values between them. An end-member composition is
    *defined* by extrapolation to zero magnesium — that zero is the most
    informative number in the row. A blanket "0 means missing" rule would erase
    the defining property of 1252 samples.
    """
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        number = float(value)
        # A bare NaN survives Python's round-trip and only explodes at the
        # Postgres jsonb cast. Coerce here, where the cause is still visible.
        return None if math.isnan(number) or math.isinf(number) else number
    return None


def _coordinate_status(lat: float | None, lon: float | None) -> str:
    """Why a sample can or cannot be placed on a map.

    ⭐ Derived, and deliberately so: it describes the SOURCE, not the sample, in
    the same spirit as `date_precision` elsewhere in this codebase. Without it
    "no marker" is indistinguishable from "we dropped the row".

    `out_of_range` is not hypothetical. All 39 Guaymas Basin samples carry
    `Latitude = 111.4`, which no latitude can be: the two axes are transposed in
    the source and the longitude's minus sign is gone. Guaymas Basin is at
    27.01 N / -111.40 W. ⛔ We do NOT transpose them back — the row keeps the
    source's numbers and simply gets no geometry.
    """
    if lat is None or lon is None:
        return "missing"
    if abs(lat) > 90 or abs(lon) > 180:
        return "out_of_range"
    return "ok"


def parse_marhys(blob: bytes) -> dict[str, Any]:
    """Parse the workbook. Returns `{"records": [...], "param_units": {...}}`.

    `param_units` maps every `params` key to the unit the source declared for
    it, so no consumer ever has to guess one from a column name.
    """
    book = openpyxl.load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    if SHEET not in book.sheetnames:
        raise MarhysFormatError(
            f"expected a sheet named {SHEET!r}, found {book.sheetnames!r}"
        )
    sheet = book[SHEET]

    rows = sheet.iter_rows(min_row=HEADER_ROW, values_only=True)
    try:
        header = next(rows)
        units = next(rows)
    except StopIteration as exc:                                # pragma: no cover
        raise MarhysFormatError("workbook has no header/units rows") from exc

    _check_layout(header, units)

    param_units = {
        str(header[i]).strip(): str(units[i] or "").strip()
        for i in range(len(header))
        if header[i] is not None and i not in _NON_PARAM_COLUMNS
    }

    records: list[dict[str, Any]] = []
    for position, row in enumerate(rows, start=1):
        # ⛔ INCLUSION IS DECIDED ON THE RAW CELL, NOT THE CLEANED ONE. An earlier
        # draft skipped rows whose cleaned Sample-ID was None — which also skips
        # every row whose id the source spells "not given". Measured: that
        # dropped 71 real samples, 71 of them with coordinates and chemistry,
        # and the row count came out at 6717 instead of 6788. A sample the
        # source did not name is still a sample.
        if not row or row[_ID] is None or str(row[_ID]).strip() == "":
            continue

        record: dict[str, Any] = {
            # ⛔ THE KEY IS THE ROW'S POSITION, NOT `Sample-ID`. Measured on
            # 2026-09-22: 6788 rows carry only 6108 distinct Sample-IDs, in 503
            # duplicate groups — it is a constructed "Site-Vent-Year" label, not
            # an accession number. Keying an upsert on it would collapse 680
            # distinct samples into each other. The published file is frozen
            # behind a DOI, so its row order is a stable identity.
            "source_row": position,
        }
        for index, field in _TEXT_COLUMNS.items():
            record[field] = _clean_text(row[index]) if index < len(row) else None
        for index, (field, _unit) in _NUMERIC_COLUMNS.items():
            record[field] = _clean_number(row[index]) if index < len(row) else None

        record["coord_status"] = _coordinate_status(record["lat"], record["lon"])

        params = {}
        for index in range(len(row)):
            if index in _NON_PARAM_COLUMNS or header[index] is None:
                continue
            value = _clean_number(row[index])
            if value is not None:
                params[str(header[index]).strip()] = value
        record["params"] = params

        records.append(record)

    return {"records": records, "param_units": param_units}


def _check_layout(header: tuple, units: tuple) -> None:
    """Fail loudly if the workbook no longer matches what this parser expects.

    ⛔ A silently reshuffled column is the worst outcome available here: every
    value would still be a plausible number, just attributed to the wrong
    element. A republished file must break this parser, not quietly repopulate
    the layer with wrong chemistry.
    """
    expected_header = {
        1: "Sample-ID", 25: "Sample type", 26: "Latitude", 27: "Longitude",
        28: "Depth", 32: "T", 36: "pH", 64: "Mg", 104: "Cl",
    }
    for index, name in expected_header.items():
        found = str(header[index]).strip() if index < len(header) and header[index] else None
        if found != name:
            raise MarhysFormatError(
                f"column {index} should be {name!r}, workbook has {found!r} — "
                "the source layout changed; re-read it before trusting any value"
            )

    for index, (field, expected_unit) in _NUMERIC_COLUMNS.items():
        found = str(units[index] or "").strip() if index < len(units) else ""
        if found != expected_unit:
            raise MarhysFormatError(
                f"{field} should be in {expected_unit!r}, workbook declares "
                f"{found!r} — refusing to store values under a unit the source "
                "no longer uses"
            )


async def fetch_marhys_workbook(timeout: float = 180.0) -> bytes:
    """Download the published workbook from PANGAEA.

    ⚠️ ~5.1 MB per call, and the file is FROZEN behind a DOI — version 4.0 was
    published 2024-10-14 and cannot change. A periodic re-download would spend
    5 MB to insert nothing, the same waste that cost the Argo sync 552 MB twice
    a day. `sync_marhys` therefore seeds once and only re-fetches on `force`.
    A version 5.0 would carry a different DOI, so it is a new ingest, not a
    silent update to this one.
    """
    async with httpx.AsyncClient(timeout=timeout, follow_redirects=True) as client:
        response = await client.get(
            MARHYS_XLSX_URL,
            headers={"User-Agent": "abyssal-claims/1.0 (+https://something-rare.com)"},
        )
        response.raise_for_status()
        return response.content
